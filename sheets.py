import os
import csv
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional, Dict, Any, List
import requests

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger("sheets_service")
logger.setLevel(logging.INFO)

HEADERS = [
    "Booking ID",
    "Timestamp (UTC)",
    "Full Name",
    "Email",
    "Phone",
    "Event",
    "Ticket Tier",
    "Tickets",
    "Special Notes",
    "Status"
]

PROJECT_ROOT = Path(__file__).parent
XLSX_FILE = PROJECT_ROOT / "bookings_spreadsheet.xlsx"
CSV_FILE = PROJECT_ROOT / "bookings_spreadsheet.csv"
LOCAL_DATA_FILE = PROJECT_ROOT / "local_bookings.json"


class GoogleSheetsManager:
    def __init__(self):
        self.creds_file = os.getenv("GOOGLE_SERVICE_ACCOUNT_FILE", "credentials.json")
        self.sheet_name = os.getenv("SPREADSHEET_NAME", "Event Bookings")
        self.sheet_id = os.getenv("SPREADSHEET_ID", "").strip()
        self.webhook_url = os.getenv("GOOGLE_SHEET_WEBHOOK_URL", "").strip()
        self.client = None
        self.worksheet = None
        self._init_local_spreadsheet()
        self._init_cloud_sheets()

    def _init_local_spreadsheet(self):
        """Initializes the local Excel and CSV spreadsheets with formatted headers."""
        # 1. Initialize CSV spreadsheet if missing
        if not CSV_FILE.exists():
            with open(CSV_FILE, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(HEADERS)
            logger.info("Initialized local CSV spreadsheet at %s", CSV_FILE.name)

        # 2. Initialize XLSX spreadsheet if missing
        if OPENPYXL_AVAILABLE and not XLSX_FILE.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Event Bookings"

            # Beautiful Google Sheets style header formatting
            header_fill = PatternFill(start_color="0F9D58", end_color="0F9D58", fill_type="solid") # Google Sheets green
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style='thin', color='D0D5DD'),
                right=Side(style='thin', color='D0D5DD'),
                top=Side(style='thin', color='D0D5DD'),
                bottom=Side(style='thin', color='D0D5DD')
            )

            ws.append(HEADERS)
            ws.row_dimensions[1].height = 26

            for col_num, header in enumerate(HEADERS, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center
                cell.border = thin_border
                # Set reasonable column widths
                col_letter = openpyxl.utils.get_column_letter(col_num)
                ws.column_dimensions[col_letter].width = max(len(header) + 6, 16)

            wb.save(XLSX_FILE)
            logger.info("Initialized formatted Excel spreadsheet at %s", XLSX_FILE.name)

    def _init_cloud_sheets(self):
        """Attempts to connect to cloud Google Sheets if credentials or webhook exist."""
        if self.webhook_url:
            return

        creds_path = Path(self.creds_file)
        if not creds_path.is_absolute():
            creds_path = PROJECT_ROOT / creds_path

        if not creds_path.exists():
            return

        try:
            import gspread
            self.client = gspread.service_account(filename=str(creds_path))
            if self.sheet_id:
                spreadsheet = self.client.open_by_key(self.sheet_id)
            else:
                spreadsheet = self.client.open(self.sheet_name)
            self.worksheet = spreadsheet.sheet1
        except Exception:
            self.client = None
            self.worksheet = None

    def is_connected(self) -> bool:
        return bool(self.webhook_url) or (self.worksheet is not None)

    def get_service_account_email(self) -> Optional[str]:
        creds_path = PROJECT_ROOT / self.creds_file
        if creds_path.exists():
            try:
                with open(creds_path, "r", encoding="utf-8") as f:
                    return json.load(f).get("client_email")
            except Exception:
                return None
        return None

    def append_booking(self, booking: Dict[str, Any]) -> Dict[str, Any]:
        """
        Directly saves booking into the spreadsheet (Excel XLSX + CSV + Cloud if configured).
        """
        now_utc = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        row = [
            booking.get("booking_id", ""),
            now_utc,
            booking.get("full_name", ""),
            booking.get("email", ""),
            booking.get("phone", ""),
            booking.get("event_name", ""),
            booking.get("ticket_tier", ""),
            booking.get("ticket_count", 1),
            booking.get("notes", ""),
            "Confirmed"
        ]

        # 1. ALWAYS directly write to the local Excel Spreadsheet (.xlsx)
        if OPENPYXL_AVAILABLE:
            try:
                if not XLSX_FILE.exists():
                    self._init_local_spreadsheet()
                wb = openpyxl.load_workbook(XLSX_FILE)
                ws = wb.active
                ws.append(row)
                
                # Apply neat cell styling
                row_idx = ws.max_row
                ws.row_dimensions[row_idx].height = 22
                thin_border = Border(
                    left=Side(style='thin', color='E5E7EB'),
                    right=Side(style='thin', color='E5E7EB'),
                    top=Side(style='thin', color='E5E7EB'),
                    bottom=Side(style='thin', color='E5E7EB')
                )
                for col_idx in range(1, len(row) + 1):
                    c = ws.cell(row=row_idx, column=col_idx)
                    c.font = Font(name="Calibri", size=10)
                    c.border = thin_border
                    if col_idx in [1, 2, 8, 10]:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")
                
                wb.save(XLSX_FILE)
            except Exception as e:
                logger.error("Error writing to Excel spreadsheet: %s", e)

        # 2. ALWAYS directly write to the local CSV Spreadsheet (.csv)
        try:
            with open(CSV_FILE, "a", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(row)
        except Exception as e:
            logger.error("Error writing to CSV spreadsheet: %s", e)

        # 3. Save to JSON store for fast browser table loading
        local_records = []
        if LOCAL_DATA_FILE.exists():
            try:
                with open(LOCAL_DATA_FILE, "r", encoding="utf-8") as f:
                    local_records = json.load(f)
            except Exception:
                local_records = []
        record = {header: val for header, val in zip(HEADERS, row)}
        local_records.append(record)
        with open(LOCAL_DATA_FILE, "w", encoding="utf-8") as f:
            json.dump(local_records, f, indent=2)

        # 4. If Cloud Webhook URL or Google Service Account is configured, mirror it there as well
        cloud_synced = False
        if self.webhook_url:
            try:
                payload = {
                    "booking_id": booking.get("booking_id", ""),
                    "timestamp": now_utc,
                    "full_name": booking.get("full_name", ""),
                    "email": booking.get("email", ""),
                    "phone": booking.get("phone", ""),
                    "event_name": booking.get("event_name", ""),
                    "ticket_tier": booking.get("ticket_tier", ""),
                    "ticket_count": booking.get("ticket_count", 1),
                    "notes": booking.get("notes", ""),
                    "status": "Confirmed",
                    "row": row
                }
                resp = requests.post(self.webhook_url, json=payload, timeout=8)
                if resp.status_code in (200, 302):
                    cloud_synced = True
            except Exception as e:
                logger.warning("Cloud webhook failed: %s", e)

        if self.worksheet is not None and not cloud_synced:
            try:
                self.worksheet.append_row(row)
                cloud_synced = True
            except Exception as e:
                logger.warning("Cloud service account failed: %s", e)

        return {
            "success": True,
            "destination": "google_spreadsheet" if cloud_synced else "spreadsheet_direct",
            "spreadsheet_file": str(XLSX_FILE.name),
            "booking_id": booking.get("booking_id"),
            "timestamp": now_utc,
            "row": row
        }

    def get_all_rows(self) -> List[Dict[str, Any]]:
        """Returns all rows in the spreadsheet."""
        if LOCAL_DATA_FILE.exists():
            try:
                with open(LOCAL_DATA_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass
        return []

    def get_recent_bookings(self, limit: int = 15) -> List[Dict[str, Any]]:
        records = self.get_all_rows()
        return records[-limit:]
